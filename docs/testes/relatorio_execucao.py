"""Relatório de execução dos testes do SIGA UEA (DoD das histórias S02 e S03).

Cruza os casos de teste da matriz (docs/matriz-ct-testes.md) com o que foi executado e
gera docs/testes/Relatorio_Execucao_Testes_SIGA_UEA.xlsx. As evidências ficam em
docs/testes/evidencias/:

  junit-sqlite.xml, junit-postgresql.xml   suíte da API (pytest --junitxml)
  e2e-*.log                                testes no navegador (docs/testes/e2e/*.mjs)
  carga-rnf01.log                          teste de carga (backend/scripts/carga_rnf01.py)
  telas/*.png                              capturas de tela dos testes no navegador

Como refazer a execução (PowerShell, a partir da raiz do projeto):

  # 1. API: SQLite e PostgreSQL (banco de teste vazio)
  cd backend
  pytest -q --junitxml=../docs/testes/evidencias/junit-sqlite.xml
  $env:TEST_DATABASE_URL="postgresql+psycopg://..."; pytest -q --junitxml=../docs/testes/evidencias/junit-postgresql.xml
  # 2. Navegador: banco de demonstração, API sem SMTP e site apontando para ela
  #    npm i --no-save playwright-core@1.48.2 (uma vez; usa o Chrome instalado)
  #    node docs/testes/e2e/cadastro_tcc.mjs, minhas_turmas.mjs, salvar_rascunho.mjs e complementares.mjs
  #    (argumentos no começo de cada arquivo), com a saída em evidencias/e2e-<nome>.log
  # 3. Carga: backend/scripts/carga_rnf01.py com a saída em evidencias/carga-rnf01.log
  # 4. python docs/testes/relatorio_execucao.py      (precisa de: pip install openpyxl)
"""

from __future__ import annotations

import re
import subprocess
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent
EVID = AQUI / "evidencias"
SAIDA = AQUI / "Relatorio_Execucao_Testes_SIGA_UEA.xlsx"
MATRIZ = RAIZ / "docs" / "matriz-ct-testes.md"

# Checagens dos roteiros de navegador que não citam o CT no texto: CT que cada uma evidencia.
# (O roteiro complementares.mjs já escreve o CT entre colchetes.)
E2E_CT = {
    "protocolo único gerado": ["CT-03.04-07", "CT-03.03-08", "CT-03.01-08", "CT-03.04-01"],
    "rascunho salvo ao continuar": ["CT-03.01-04"],
    "orientador aceitou a orientação": ["CT-06.01-01"],
    "envio bloqueado enquanto há pendências": ["CT-06.01-03"],
    "histórico registra a submissão": ["CT-03.03-01"],
    "sem botão de novo cadastro: um TCC por período": ["CT-03.04-04"],
    "coordenação cadastra a aluna": ["CT-01.03-01"],
    "aluna ativa a conta e volta para o login": ["CT-01.05-02", "CT-01.04-01"],
    "Salvar rascunho mantém a aluna no cadastro": ["CT-03.01-04"],
    "rascunho aparece na lista, sem protocolo": ["CT-03.04-02"],
}
ROTEIROS = {
    "cadastro_tcc": "Cadastro de TCC de ponta a ponta (aluno → orientador → aluno)",
    "minhas_turmas": "Minhas turmas (professor, aluno e coordenação)",
    "salvar_rascunho": "Cadastro de aluna, ativação e Salvar rascunho",
    "complementares": "Casos de tela e responsividade (360 e 768 px)",
}
EVID_TELA = {  # captura de tela que acompanha cada CT executado no navegador
    "CT-01.06-09": "telas/complementar-login-360.png; telas/complementar-login-768.png",
    "CT-01.06-06": "telas/complementar-login-360.png",
    "CT-01.01-08": "telas/complementar-docente-360.png; telas/complementar-docente-768.png",
    "CT-01.02-01": "telas/complementar-docente-768.png",
    "CT-01.09-06": "telas/complementar-perfil-aluno-360.png; telas/complementar-perfil-aluno-768.png",
    "CT-01.10-03": "telas/complementar-perfil-banca.png",
    "CT-03.01-01": "telas/complementar-cadastro-artigo-360.png",
    "CT-03.01-07": "telas/complementar-cadastro-artigo-360.png; telas/complementar-cadastro-768.png",
    "CT-03.02-01": "telas/complementar-cadastro-artigo-360.png",
    "CT-03.02-02": "telas/complementar-cadastro-monografia-360.png",
    "CT-03.02-07": "telas/complementar-cadastro-monografia-360.png",
    "CT-03.04-07": "telas/cadastro-5-protocolo.png; telas/cadastro-6-acompanhamento.png",
    "CT-03.03-08": "telas/cadastro-5-protocolo.png",
    "CT-03.01-08": "telas/cadastro-1-trabalho.png; telas/cadastro-5-protocolo.png",
    "CT-03.01-04": "telas/rascunho-1-salvo-sem-sair.png",
    "CT-06.01-01": "telas/cadastro-4-orientacoes.png",
    "CT-06.01-03": "telas/cadastro-3-revisar-pendente.png",
}

# ---------------------------------------------------------------- leitura


def ler_matriz():
    cts, rbn = [], None
    for linha in MATRIZ.read_text(encoding="utf8").splitlines():
        if linha.startswith("## "):
            rbn = linha[3:].strip()
            continue
        if not linha.startswith("| CT-"):
            continue
        c = [x.strip() for x in linha.strip().strip("|").split("|")]
        testes = [] if c[4] in ("—", "") else [t.strip() for t in c[4].split("<br>")]
        cts.append(dict(rbn=rbn, ct=c[0], tipo=c[1], passos=c[2], situacao=c[3], testes=testes,
                        obs="" if c[5] == "—" else c[5]))
    return cts


def ler_junit(arq: Path):
    """{arquivo.py::funcao: [(parametros, resultado, tempo)]}; resultado: Aprovado, Pulado ou Reprovado."""
    raiz = ET.parse(arq).getroot()
    suite = raiz.find("testsuite")
    out = defaultdict(list)
    for tc in suite.iter("testcase"):
        modulo = tc.get("classname").split(".")[-1] + ".py"
        nome = tc.get("name")
        base, _, param = nome.partition("[")
        if tc.find("failure") is not None or tc.find("error") is not None:
            res = "Reprovado"
        elif tc.find("skipped") is not None:
            res = "Pulado"
        else:
            res = "Aprovado"
        out[f"{modulo}::{base}"].append((param.rstrip("]"), res, float(tc.get("time", 0))))
    info = dict(testes=int(suite.get("tests")), falhas=int(suite.get("failures")) + int(suite.get("errors")),
                pulados=int(suite.get("skipped")), tempo=float(suite.get("time")),
                quando=datetime.fromisoformat(suite.get("timestamp")))
    return out, info


def ler_e2e():
    """Lista de checagens: (roteiro, texto, resultado, [CTs])."""
    out = []
    for arq in sorted(EVID.glob("e2e-*.log")):
        roteiro = arq.stem.removeprefix("e2e-")
        for linha in arq.read_text(encoding="utf8").splitlines():
            m = re.match(r"^(OK|FALHA)\s+(.*)$", linha)
            if not m:
                continue
            texto = m.group(2).strip()
            cts = re.findall(r"\[(CT-[\d.\-]+)\]", texto)
            texto = re.sub(r"^\[CT-[\d.\-]+\]\s*", "", texto)
            if not cts:
                cts = next((v for k, v in E2E_CT.items() if texto.startswith(k)), [])
            out.append((roteiro, texto, "Aprovado" if m.group(1) == "OK" else "Reprovado", cts))
    return out


def ler_carga():
    txt = (EVID / "carga-rnf01.log").read_text(encoding="utf8")
    num = lambda r: float(re.search(r, txt).group(1))  # noqa: E731
    return dict(texto=txt, media=num(r"média ([\d.]+) s"), p95=num(r"p95 ([\d.]+) s"), maximo=num(r"máximo ([\d.]+) s"),
                distintos=int(num(r"distintos: (\d+)")), ok="RNF01 OK" in txt and "CT-03.04-05 OK" in txt,
                quando=re.search(r"Execução: ([^·]+)", txt).group(1).strip())


def git(*args):
    return subprocess.run(["git", *args], cwd=RAIZ, capture_output=True, text=True, encoding="utf8").stdout.strip()


# ---------------------------------------------------------------- consolidação


def consolidar(cts, sqlite, pg, e2e, carga):
    por_ct = defaultdict(list)
    for roteiro, texto, res, lista in e2e:
        for ct in lista:
            por_ct[ct].append((roteiro, texto, res))
    linhas = []
    for c in cts:
        res_sqlite, res_pg = [], []
        for t in c["testes"]:
            res_sqlite += [r for _, r, _ in sqlite.get(t, [("", "Não encontrado", 0)])]
            res_pg += [r for _, r, _ in pg.get(t, [("", "Não encontrado", 0)])]
        nav = por_ct.get(c["ct"], [])
        carga_ct = c["ct"] in ("CT-03.03-07", "CT-03.04-05")
        formas = []
        if c["testes"]:
            formas.append("API (pytest)")
        if nav:
            formas.append("Navegador")
        if carga_ct:
            formas.append("Carga")

        def resumo(lista):
            if not lista:
                return "—"
            if "Reprovado" in lista or "Não encontrado" in lista:
                return "Reprovado" if "Reprovado" in lista else "Não encontrado"
            return "Pulado" if all(r == "Pulado" for r in lista) else "Aprovado"

        r_sqlite, r_pg = resumo(res_sqlite), resumo(res_pg)
        r_nav = resumo([r for _, _, r in nav])
        r_carga = ("Aprovado" if carga["ok"] else "Reprovado") if carga_ct else "—"
        todos = [r for r in (r_sqlite, r_pg, r_nav, r_carga) if r != "—"]
        if any(r in ("Reprovado", "Não encontrado") for r in todos):
            final = "Reprovado"
        elif c["situacao"] == "Pendente":
            final = "Pendente"
        elif not todos or all(r == "Pulado" for r in todos):
            final = "Não executado"
        else:
            final = "Aprovado"

        obs = c["obs"]
        if c["situacao"] in ("Manual", "Parcial") and nav:
            obs = (obs + " " if obs else "") + "Parte de tela executada no navegador nesta rodada."
        if carga_ct:
            p95 = f"{carga['p95']:.2f}".replace(".", ",")
            obs = (obs + " " if obs else "") + (f"30 submissões simultâneas: p95 {p95} s (meta 3 s), "
                                                 f"{carga['distintos']} protocolos distintos.")
        evid = []
        if c["testes"]:
            evid.append("junit-sqlite.xml; junit-postgresql.xml")
        evid += sorted({f"e2e-{rot}.log" for rot, _, _ in nav})
        if carga_ct:
            evid.append("carga-rnf01.log")
        if c["ct"] in EVID_TELA:
            evid.append(EVID_TELA[c["ct"]])
        linhas.append(dict(c, formas=" + ".join(formas) or "—", r_sqlite=r_sqlite, r_pg=r_pg, r_nav=r_nav,
                           r_carga=r_carga, final=final, obs_final=obs, evid="; ".join(evid) or "—",
                           checagens=" | ".join(t for _, t, _ in nav)))
    return linhas


# ---------------------------------------------------------------- planilha

FONTE = "Arial"
AZUL = "1F4E79"
CORES = {"Aprovado": ("E2F0D9", "375623"), "Pendente": ("FFF2CC", "7F6000"), "Reprovado": ("F8CBAD", "9C0006"),
         "Não executado": ("EDEDED", "3A3A3A"), "Pulado": ("FFF2CC", "7F6000")}
FINO = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def f(bold=False, cor="000000", tam=10, it=False):
    return Font(name=FONTE, size=tam, bold=bold, color=cor, italic=it)


def tabela(ws, linha0, cabecalho, linhas, larguras):
    for j, (titulo, w) in enumerate(zip(cabecalho, larguras), start=1):
        c = ws.cell(row=linha0, column=j, value=titulo)
        c.font = f(True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDA
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, valores in enumerate(linhas, start=linha0 + 1):
        for j, v in enumerate(valores, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = f()
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDA
    ultima = linha0 + len(linhas)
    ws.auto_filter.ref = f"A{linha0}:{get_column_letter(len(cabecalho))}{ultima}"
    ws.freeze_panes = ws.cell(row=linha0 + 1, column=1)
    return ultima


def colorir(ws, coluna, linha0, ultima):
    """Cor do resultado pela regra condicional: acompanha se alguém editar a célula."""
    faixa = f"{coluna}{linha0 + 1}:{coluna}{ultima}"
    for valor, (fundo, texto) in CORES.items():
        ws.conditional_formatting.add(faixa, FormulaRule(
            formula=[f'${coluna}{linha0 + 1}="{valor}"'], fill=PatternFill("solid", fgColor=fundo),
            font=Font(name=FONTE, color=texto, bold=True)))


def montar(linhas, sqlite_info, pg_info, junit_sqlite, junit_pg, e2e, carga):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True  # o Excel e o Google Planilhas calculam os totais ao abrir

    # ---------------- Execução por CT
    ws = wb.active
    ws.title = "Execucao por CT"
    cab = ["Requisito", "CT", "Tipo", "Passos", "Situação na matriz", "Como foi executado", "Teste(s) automatizado(s)",
           "API · SQLite", "API · PostgreSQL 16", "Navegador", "Carga", "Resultado final", "Evidência",
           "Checagens no navegador", "Observação"]
    dados = [[l["rbn"], l["ct"], l["tipo"], l["passos"], l["situacao"], l["formas"], "\n".join(l["testes"]) or "—",
              l["r_sqlite"], l["r_pg"], l["r_nav"], l["r_carga"], l["final"], l["evid"], l["checagens"] or "—",
              l["obs_final"] or "—"] for l in linhas]
    ult = tabela(ws, 1, cab, dados, [11, 13, 12, 30, 12, 16, 48, 11, 12, 11, 10, 14, 34, 44, 44])
    for col in "HIJKL":
        colorir(ws, col, 1, ult)
    n_ct = len(linhas)

    # ---------------- Testes automatizados
    wt = wb.create_sheet("Testes automatizados")
    todos = sorted(set(junit_sqlite) | set(junit_pg))
    ct_de = defaultdict(list)
    for l in linhas:
        for t in l["testes"]:
            ct_de[t].append(l["ct"])
    dados_t = []
    for t in todos:
        s, p = junit_sqlite.get(t, []), junit_pg.get(t, [])
        res = lambda xs: ("Reprovado" if any(r == "Reprovado" for _, r, _ in xs) else  # noqa: E731
                          "Pulado" if xs and all(r == "Pulado" for _, r, _ in xs) else "Aprovado" if xs else "—")
        arq, _, func = t.partition("::")
        dados_t.append([arq, func, len(s) or len(p), ", ".join(ct_de.get(t, [])) or "complementar (sem CT na matriz)",
                        res(s), round(sum(x[2] for x in s), 3), res(p), round(sum(x[2] for x in p), 3)])
    ult_t = tabela(wt, 1, ["Arquivo", "Teste", "Execuções (parâmetros)", "CT(s)", "SQLite", "Tempo SQLite (s)",
                           "PostgreSQL 16", "Tempo PostgreSQL (s)"], dados_t, [24, 70, 12, 26, 12, 12, 13, 13])
    colorir(wt, "E", 1, ult_t)
    colorir(wt, "G", 1, ult_t)

    # ---------------- Testes no navegador
    wn = wb.create_sheet("Testes no navegador")
    dados_n = [[ROTEIROS.get(rot, rot), f"docs/testes/e2e/{rot}.mjs", txt, ", ".join(cts) or "—", res,
                f"e2e-{rot}.log"] for rot, txt, res, cts in e2e]
    ult_n = tabela(wn, 1, ["Roteiro", "Arquivo", "Checagem", "CT(s)", "Resultado", "Evidência"], dados_n,
                   [40, 32, 70, 22, 12, 24])
    colorir(wn, "E", 1, ult_n)

    # ---------------- Desempenho
    wd = wb.create_sheet("Desempenho RNF01")
    wd["A1"] = "Teste de carga do cadastro de TCC (RNF01, CT-03.03-07 e CT-03.04-05)"
    wd["A1"].font = f(True, AZUL, 13)
    linhas_d = [
        ("Data da execução", carga["quando"], "carga-rnf01.log"),
        ("Submissões simultâneas", 30, "backend/scripts/carga_rnf01.py"),
        ("Tempo médio (s)", carga["media"], "medido do envio até o retorno do protocolo"),
        ("p95 (s)", carga["p95"], "percentil 95 dos 30 tempos"),
        ("Tempo máximo (s)", carga["maximo"], ""),
        ("Meta do RNF01: p95 até (s)", 3, "planilha de requisitos, RNF01"),
        ("Protocolos distintos", carga["distintos"], "RNF06: nenhum protocolo repetido"),
    ]
    for i, (k, v, nota) in enumerate(linhas_d, start=3):
        wd.cell(row=i, column=1, value=k).font = f(True)
        c = wd.cell(row=i, column=2, value=v)
        c.font = f(cor="0000FF")
        if isinstance(v, float):
            c.number_format = "0.000"
        wd.cell(row=i, column=3, value=nota).font = f(it=True, cor="595959")
    wd["A11"] = "Resultado RNF01 (p95 abaixo da meta)"
    wd["B11"] = '=IF(B6<B8,"Aprovado","Reprovado")'
    wd["A12"] = "Resultado RNF06 (30 protocolos distintos)"
    wd["B12"] = '=IF(B9=B4,"Aprovado","Reprovado")'
    for r in (11, 12):
        wd[f"A{r}"].font = f(True)
        wd[f"B{r}"].font = f(True)
    wd["A14"] = "Valores em azul: copiados da saída do teste (evidencias/carga-rnf01.log)."
    wd["A14"].font = f(it=True, cor="595959")
    wd.column_dimensions["A"].width = 40
    wd.column_dimensions["B"].width = 22
    wd.column_dimensions["C"].width = 48
    colorir(wd, "B", 10, 12)

    # ---------------- Resumo (primeira aba)
    wr = wb.create_sheet("Resumo", 0)
    commit = git("rev-parse", "--short", "HEAD")
    wr["A1"] = "SIGA UEA · Relatório de execução de testes"
    wr["A1"].font = f(True, AZUL, 15)
    wr["A2"] = ("Histórias das sprints S02 e S03 (RBN01.01 a RBN01.10, RBN03.01 a RBN03.04 e RBN06.01) · "
                "casos de teste da matriz docs/matriz-ct-testes.md")
    wr["A2"].font = f(it=True, cor="595959")
    meta = [
        ("Data da execução", sqlite_info["quando"].strftime("%d/%m/%Y %H:%M") + " (America/Manaus)"),
        ("Versão testada (commit)", f"{commit} · branch dev · repositório SIGA-UEA/sigauea"),
        ("Ambientes", "API Python 3.12 + FastAPI 0.143 com SQLite e com PostgreSQL 16; site Next.js 16 (build de "
                      "produção) no Google Chrome; e-mail desligado (SMTP vazio) para não enviar mensagens reais"),
        ("Ferramentas", "pytest (API), Playwright com Chrome (navegador), script de carga em Python (httpx)"),
        ("Responsável pela execução", "Igor Leite de Almeida"),
    ]
    for i, (k, v) in enumerate(meta, start=4):
        wr.cell(row=i, column=1, value=k).font = f(True)
        wr.cell(row=i, column=2, value=v).font = f()
        wr.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    wr["A10"] = "Casos de teste (CT) da matriz"
    wr["A10"].font = f(True, AZUL, 12)
    rng = f"'Execucao por CT'!$L$2:$L${n_ct + 1}"
    linha = 11
    for rot in ("Aprovado", "Pendente", "Não executado", "Reprovado"):
        wr.cell(row=linha, column=1, value=rot).font = f(True)
        wr.cell(row=linha, column=2, value=f'=COUNTIF({rng},"{rot}")').font = f()
        c = wr.cell(row=linha, column=3, value=f"=IF($B$15=0,0,B{linha}/$B$15)")
        c.number_format = "0.0%"
        c.font = f()
        linha += 1
    wr.cell(row=15, column=1, value="Total de CTs").font = f(True)
    wr.cell(row=15, column=2, value="=SUM(B11:B14)").font = f(True)
    colorir(wr, "A", 10, 14)

    wr["A17"] = "Execuções"
    wr["A17"].font = f(True, AZUL, 12)
    rt = f"'Testes automatizados'"
    exec_linhas = [
        ("Testes da API no SQLite: aprovados", f'=COUNTIF({rt}!$E$2:$E${ult_t},"Aprovado")'),
        ("Testes da API no SQLite: pulados (funcionalidade futura)", f'=COUNTIF({rt}!$E$2:$E${ult_t},"Pulado")'),
        ("Testes da API no SQLite: reprovados", f'=COUNTIF({rt}!$E$2:$E${ult_t},"Reprovado")'),
        ("Testes da API no PostgreSQL 16: aprovados", f'=COUNTIF({rt}!$G$2:$G${ult_t},"Aprovado")'),
        ("Testes da API no PostgreSQL 16: pulados", f'=COUNTIF({rt}!$G$2:$G${ult_t},"Pulado")'),
        ("Testes da API no PostgreSQL 16: reprovados", f'=COUNTIF({rt}!$G$2:$G${ult_t},"Reprovado")'),
        ("Checagens no navegador: aprovadas", f"=COUNTIF('Testes no navegador'!$E$2:$E${ult_n},\"Aprovado\")"),
        ("Checagens no navegador: reprovadas", f"=COUNTIF('Testes no navegador'!$E$2:$E${ult_n},\"Reprovado\")"),
        ("Carga RNF01: p95 em segundos (meta 3 s)", "='Desempenho RNF01'!B6"),
        ("Carga RNF01: resultado", "='Desempenho RNF01'!B11"),
    ]
    for i, (k, v) in enumerate(exec_linhas, start=18):
        wr.cell(row=i, column=1, value=k).font = f(True)
        c = wr.cell(row=i, column=2, value=v)
        c.font = f(cor="008000")
        if "p95" in k:
            c.number_format = "0.000"
    n = 18 + len(exec_linhas) + 1
    wr.cell(row=n, column=1, value=(
        f"Funções de teste da API: {len(todos)} ({sqlite_info['testes']} execuções contando os parâmetros). "
        "Um teste pulado é um caso Pendente na matriz: a funcionalidade ainda não existe (CT-01.10-06 depende do "
        "módulo de ideias; CT-06.01-04, da troca de orientador pela coordenação).")).font = f(it=True, cor="595959")
    wr.cell(row=n + 1, column=1, value=(
        "Os CTs com tempo simulado (expiração de link, bloqueio de 15 min, sessão de 60 min) adiantam o relógio no "
        "banco em vez de esperar; ficam como Aprovado com essa observação.")).font = f(it=True, cor="595959")
    wr.cell(row=n + 2, column=1, value=(
        "Os totais desta aba são fórmulas sobre as outras abas e se atualizam ao abrir o arquivo; "
        "verde = valor vindo de outra aba.")).font = f(it=True, cor="595959")

    wr.cell(row=n + 4, column=1, value="Por requisito").font = f(True, AZUL, 12)
    cab_r = ["Requisito", "CTs", "Aprovados", "Pendentes", "Não executados", "Reprovados"]
    for j, t in enumerate(cab_r, start=1):
        c = wr.cell(row=n + 5, column=j, value=t)
        c.font = f(True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.border = BORDA
    rbns = list(dict.fromkeys(l["rbn"] for l in linhas))
    a = f"'Execucao por CT'!$A$2:$A${n_ct + 1}"
    for i, rbn in enumerate(rbns, start=n + 6):
        wr.cell(row=i, column=1, value=rbn)
        wr.cell(row=i, column=2, value=f'=COUNTIF({a},A{i})')
        for j, rot in enumerate(("Aprovado", "Pendente", "Não executado", "Reprovado"), start=3):
            wr.cell(row=i, column=j, value=f'=COUNTIFS({a},$A{i},{rng},"{rot}")')
        for j in range(1, 7):
            wr.cell(row=i, column=j).font = f(cor="000000" if j == 1 else "008000")
            wr.cell(row=i, column=j).border = BORDA
    wr.column_dimensions["A"].width = 52
    wr.column_dimensions["B"].width = 90
    for col in "CDEF":
        wr.column_dimensions[col].width = 14
    wr.column_dimensions["B"].width = 90

    # ---------------- Evidências
    we = wb.create_sheet("Evidencias")
    arqs = sorted(p.relative_to(EVID).as_posix() for p in EVID.rglob("*") if p.is_file())
    desc = {
        "junit-sqlite.xml": "Saída do pytest (formato JUnit) da suíte da API no SQLite",
        "junit-postgresql.xml": "Saída do pytest (formato JUnit) da suíte da API no PostgreSQL 16",
        "carga-rnf01.log": "Saída do teste de carga com 30 submissões simultâneas",
    }
    dados_e = []
    for a_ in arqs:
        if a_.startswith("e2e-"):
            d = "Saída do roteiro: " + ROTEIROS.get(a_.removeprefix("e2e-").removesuffix(".log"), "")
        elif a_.startswith("telas/"):
            d = "Captura de tela do teste no navegador"
        else:
            d = desc.get(a_, "")
        dados_e.append([f"docs/testes/evidencias/{a_}", d])
    tabela(we, 1, ["Arquivo (no repositório)", "Conteúdo"], dados_e, [70, 70])

    for w in wb.worksheets:
        w.sheet_view.showGridLines = w.title not in ("Resumo", "Desempenho RNF01")
    wb.save(SAIDA)


def main():
    cts = ler_matriz()
    junit_sqlite, sqlite_info = ler_junit(EVID / "junit-sqlite.xml")
    junit_pg, pg_info = ler_junit(EVID / "junit-postgresql.xml")
    e2e = ler_e2e()
    carga = ler_carga()
    linhas = consolidar(cts, junit_sqlite, junit_pg, e2e, carga)
    montar(linhas, sqlite_info, pg_info, junit_sqlite, junit_pg, e2e, carga)

    # Conferência em Python dos mesmos totais que a planilha calcula por fórmula.
    finais = defaultdict(int)
    for l in linhas:
        finais[l["final"]] += 1
    print(f"{len(linhas)} CTs:", dict(finais))
    print(f"API SQLite: {sqlite_info['testes']} execuções, {sqlite_info['falhas']} falhas, {sqlite_info['pulados']} pulados")
    print(f"API PostgreSQL: {pg_info['testes']} execuções, {pg_info['falhas']} falhas, {pg_info['pulados']} pulados")
    print(f"navegador: {len(e2e)} checagens, {sum(r == 'Reprovado' for _, _, r, _ in e2e)} reprovadas")
    print(f"carga: p95 {carga['p95']} s, {carga['distintos']} protocolos distintos")
    sem_teste = [t for l in linhas for t in l["testes"] if t not in junit_sqlite]
    print("testes da matriz não encontrados na execução:", sem_teste or "nenhum")
    print("CTs da matriz sem nenhuma execução:", [l["ct"] for l in linhas if l["final"] == "Não executado"] or "nenhum")
    print("arquivo:", SAIDA)


if __name__ == "__main__":
    main()
