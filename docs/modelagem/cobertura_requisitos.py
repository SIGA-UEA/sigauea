"""Confere se toda RBN e RN da planilha de requisitos aparece no modelo lógico.

Cada requisito tem de estar citado em alguma tabela (campo requisitos) ou em
SEM_TABELA, com o motivo. Sai com erro se faltar algum.

    pip install openpyxl
    python docs/modelagem/cobertura_requisitos.py ["Requisitos_SIGA_UEA (4).xlsx"]
"""

import re
import sys
from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
import modelo_logico as ml  # noqa: E402

PLANILHA = AQUI.parent.parent / "Requisitos_SIGA_UEA (4).xlsx"


def citados(texto: str) -> set[str]:
    """Os identificadores do texto, abrindo intervalos como "RBN03.01 a RBN03.05"."""
    out = set(re.findall(r"R[A-Z]{0,3}\d{2}(?:\.\d{2})?", texto))
    for ini, fim in re.findall(r"(RBN\d{2}\.\d{2}) a (RBN\d{2}\.\d{2})", texto):
        out |= {f"{ini[:5]}.{n:02d}" for n in range(int(ini[6:]), int(fim[6:]) + 1)}
    return out


def main(planilha: Path) -> int:
    no_modelo = set()
    for t in ml.T:
        no_modelo |= citados(t.requisitos)
    no_modelo |= citados(" ".join(ml.SEM_TABELA))

    wb = openpyxl.load_workbook(planilha, read_only=True)
    requisitos = [r[2].split(" - ")[0] for r in wb["Matriz"].iter_rows(min_row=2, values_only=True)
                  if len(r) > 2 and r[2] and str(r[2]).startswith("RBN")]
    requisitos += [r[0] for r in wb["RN"].iter_rows(min_row=2, values_only=True) if r and r[0]]

    faltando = [r for r in requisitos if r not in no_modelo]
    print(f"{len(requisitos)} requisitos na planilha; {len(requisitos) - len(faltando)} cobertos pelo modelo")
    for r in faltando:
        print("  sem cobertura:", r)
    return 1 if faltando else 0


if __name__ == "__main__":
    raise SystemExit(main(Path(sys.argv[1]) if len(sys.argv) > 1 else PLANILHA))
